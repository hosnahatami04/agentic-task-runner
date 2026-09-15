"""Reads text out of PDF, Word, and Excel files inside the sandboxed workspace."""

from __future__ import annotations

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from agent.models import Observation

from .registry import tool
from .sandbox import resolve_in_sandbox

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx"}


def _read_pdf(path) -> str:
    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip()


def _read_docx(path) -> str:
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    return "\n".join(paragraphs).strip()


def _read_xlsx(path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"[{sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                lines.append(", ".join(cells))
    return "\n".join(lines).strip()


@tool(
    name="read_document",
    description=(
        "Reads text content from a PDF (.pdf), Word (.docx), or Excel (.xlsx) "
        "file inside the workspace."
    ),
    parameters={
        "type": "object",
        "properties": {"path": {"type": "string"}},
        "required": ["path"],
    },
)
def read_document(path: str) -> Observation:
    try:
        full_path = resolve_in_sandbox(path)
    except ValueError as exc:
        return Observation(success=False, error=str(exc))

    if not full_path.is_file():
        return Observation(success=False, error=f"File not found: {path}")

    extension = full_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        return Observation(
            success=False,
            error=f"Unsupported file type '{extension}'. Supported: pdf, docx, xlsx.",
        )

    try:
        if extension == ".pdf":
            content = _read_pdf(full_path)
        elif extension == ".docx":
            content = _read_docx(full_path)
        else:
            content = _read_xlsx(full_path)
    except Exception as exc:
        return Observation(success=False, error=f"Could not read '{path}': {exc}")

    if not content:
        return Observation(success=False, error=f"No readable text found in '{path}'")

    return Observation(success=True, data=content)
