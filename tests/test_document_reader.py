"""Tests for src/tools/document_reader.py."""

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from tools import document_reader, sandbox


@pytest.fixture(autouse=True)
def _use_tmp_workspace(tmp_path, monkeypatch):
    monkeypatch.setattr(sandbox, "WORKSPACE_ROOT", tmp_path)


def _make_blank_pdf(path):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(path, "wb") as f:
        writer.write(f)
    return path


def _make_docx(path, text="Hello from a Word document"):
    doc = Document()
    doc.add_paragraph(text)
    doc.save(path)
    return path


def _make_xlsx(path, rows=(("name", "amount"), ("widgets", 42))):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


def test_reads_docx_text(tmp_path):
    _make_docx(tmp_path / "notes.docx", "Hello from a Word document")

    result = document_reader.read_document("notes.docx")

    assert result.success is True
    assert "Hello from a Word document" in result.data


def test_reads_xlsx_rows(tmp_path):
    _make_xlsx(tmp_path / "data.xlsx")

    result = document_reader.read_document("data.xlsx")

    assert result.success is True
    assert "widgets" in result.data
    assert "42" in result.data


def test_blank_pdf_has_no_extractable_text(tmp_path):
    _make_blank_pdf(tmp_path / "blank.pdf")

    result = document_reader.read_document("blank.pdf")

    assert result.success is False
    assert "no readable text" in result.error.lower()


def test_missing_file_returns_error():
    result = document_reader.read_document("does_not_exist.pdf")

    assert result.success is False
    assert "not found" in result.error.lower()


def test_unsupported_extension_returns_error(tmp_path):
    (tmp_path / "notes.txt").write_text("plain text")

    result = document_reader.read_document("notes.txt")

    assert result.success is False
    assert "unsupported" in result.error.lower()


def test_rejects_path_traversal():
    result = document_reader.read_document("../../secret.pdf")

    assert result.success is False
    assert "sandbox" in result.error.lower()
