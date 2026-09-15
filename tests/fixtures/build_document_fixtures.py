"""Builds small, realistic PDF/Word/Excel task fixtures.

Run directly to (re)generate the files in tests/fixtures/:
    python tests/fixtures/build_document_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook

FIXTURES_DIR = Path(__file__).parent


def _build_simple_pdf(lines: list[str]) -> bytes:
    """Builds a minimal single-page PDF with real, extractable text.

    pypdf can read PDFs but can't draw text onto one without an extra
    rendering library, so the PDF content stream is written by hand here
    using the plain-text PDF syntax (Helvetica is a standard PDF font,
    no embedding needed).
    """
    text_ops = ["BT", "/F1 14 Tf", "50 740 Td"]
    for i, line in enumerate(lines):
        if i > 0:
            text_ops.append("0 -22 Td")
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        text_ops.append(f"({escaped}) Tj")
    text_ops.append("ET")
    content = "\n".join(text_ops).encode()

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 612 792] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"

    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode()

    return bytes(out)


def build_policy_pdf() -> None:
    pdf_bytes = _build_simple_pdf(
        [
            "Company Policy Handbook",
            "All employees are entitled to 18 paid vacation days per year.",
        ]
    )
    (FIXTURES_DIR / "policy.pdf").write_bytes(pdf_bytes)


def build_team_notes_docx() -> None:
    doc = Document()
    doc.add_paragraph("Team Notes - Q1 Planning")
    doc.add_paragraph(
        "The Engineering team will onboard 3 new hires in March. "
        "Budget approved for the new hires is 210000 dollars total."
    )
    doc.add_paragraph("Marketing requested a review of the Q2 campaign timeline.")
    doc.save(FIXTURES_DIR / "team_notes.docx")


def build_expenses_xlsx() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Expenses"
    sheet.append(["Category", "Amount"])
    sheet.append(["Travel", 4200])
    sheet.append(["Software", 1800])
    sheet.append(["Office Supplies", 650])
    workbook.save(FIXTURES_DIR / "expenses.xlsx")


if __name__ == "__main__":
    build_policy_pdf()
    build_team_notes_docx()
    build_expenses_xlsx()
    print("Wrote policy.pdf, team_notes.docx, expenses.xlsx")
